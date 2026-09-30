"""Convert the TPXimpact Progression Framework spreadsheets into framework.json.
Usage: python3 extract_framework.py <DT billable skills.xlsx> <behaviours and impact matrix.xlsx> <out.json>
"""
import json, re, sys
import openpyxl

SKILL_LEVELS = ["Learner", "Contributor", "Skilled", "Expert", "Leader", "Driver"]
GRADES = [
    {"code": "G6", "name": "Graduate", "number": 6},
    {"code": "G7", "name": "Junior", "number": 7},
    {"code": "G8", "name": "Mid", "number": 8},
    {"code": "G9", "name": "Senior", "number": 9},
    {"code": "G10", "name": "Lead", "number": 10},
    {"code": "G11", "name": "Principal", "number": 11},
    {"code": "G12", "name": "Head of", "number": 12},
]

def clean(t):
    if t is None:
        return ""
    t = str(t).replace("\r", "")
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()

def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")

def main(skills_path, behav_path, out_path):
    wb = openpyxl.load_workbook(skills_path, data_only=True)
    defs = {}
    for r in wb["Skill definitions"].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        name = clean(r[0])
        defs[name.lower()] = {
            "definition": clean(r[1]),
            "levels": {lvl: clean(r[2 + i]) for i, lvl in enumerate(SKILL_LEVELS)},
        }

    roles = {}
    skills = {}
    for r in wb["Skill by role (with description"].iter_rows(min_row=2, values_only=True):
        if not r[2] or not r[3]:
            continue
        cap, practice, role, skill = (clean(x) for x in r[:4])
        rid = slug(role)
        sid = slug(skill)
        roleobj = roles.setdefault(rid, {"id": rid, "name": role, "capability": cap,
                                         "practice": practice, "skills": []})
        expected = {}
        for gi, cell in enumerate(r[4:10]):
            txt = clean(cell)
            grade = GRADES[gi]["code"]
            level = None
            desc = ""
            m = re.match(r"^(Learner|Contributor|Skilled|Expert|Leader)\s*:?\s*(.*)$", txt, re.S)
            if m:
                level = m.group(1)
                desc = m.group(2).strip()
            expected[grade] = level
            if level:
                s = skills.setdefault(sid, {"id": sid, "name": skill, "definition": "", "levels": {}})
                if desc and not s["levels"].get(level):
                    s["levels"][level] = desc
        s = skills.setdefault(sid, {"id": sid, "name": skill, "definition": "", "levels": {}})
        d = defs.get(skill.lower())
        if d:
            s["definition"] = s["definition"] or d["definition"]
            for lvl, txt in d["levels"].items():
                if txt and txt.lower() != "n/a" and not s["levels"].get(lvl):
                    s["levels"][lvl] = txt
        roleobj["skills"].append({"skillId": sid, "expected": expected})

    # Behaviours (v3) and Impact (per-level descriptors from V2 Impact, definitions from v3 Impact)
    wb2 = openpyxl.load_workbook(behav_path, data_only=True)
    behaviours = []
    for r in wb2["Behaviours"].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        name = clean(r[0])
        # Columns: Junior/Graduate(6/7) | Mid(8) | Senior(9) | Lead(10) | Principal(11) | Head of(12)
        cols = [clean(c) for c in r[2:8]]
        levels = {"G6": cols[0], "G7": cols[0], "G8": cols[1], "G9": cols[2],
                  "G10": cols[3], "G11": cols[4], "G12": cols[5]}
        behaviours.append({"id": slug(name), "name": name, "type": "BEHAVIOUR",
                           "definition": clean(r[1]), "levels": levels})

    impact_areas = []
    for r in wb2["Impact"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            impact_areas.append({"name": clean(r[0]), "definition": clean(r[1])})

    impacts = []
    for r in wb2["V2 Impact"].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        name = clean(r[0]).replace("\n", " ")
        name = re.sub(r"\s+", " ", name).strip()
        if name.lower() == "overview":
            continue
        cols = [clean(c) for c in r[1:7]]
        # Junior/Graduate | Mid | Senior | Lead | Principal | Head of
        levels = {"G6": cols[0], "G7": cols[0], "G8": cols[1], "G9": cols[2],
                  "G10": cols[3], "G11": cols[4], "G12": cols[5]}
        levels = {k: ("" if v.upper() == "N/A" else v) for k, v in levels.items()}
        impacts.append({"id": slug(name), "name": name, "type": "IMPACT",
                        "definition": "", "levels": levels})

    out = {
        "source": {
            "skills": "Progression Framework - DT billable skills.xlsx (v1.0, Progression Assessment 2026)",
            "behaviours": "Progression Framework - behaviours and impact matrix.xlsx (v3.0)",
        },
        "grades": GRADES,
        "skillLevels": SKILL_LEVELS[:5],
        "roles": sorted(roles.values(), key=lambda x: (x["capability"], x["practice"], x["name"])),
        "skills": sorted(skills.values(), key=lambda x: x["name"]),
        "behaviours": behaviours,
        "impactAreas": impact_areas,
        "impacts": impacts,
    }
    with open(out_path, "w") as f:
        json.dump(out, f, indent=1, ensure_ascii=False)
    print(f"roles={len(out['roles'])} skills={len(out['skills'])} behaviours={len(behaviours)} impacts={len(impacts)}")

if __name__ == "__main__":
    main(*sys.argv[1:4])
